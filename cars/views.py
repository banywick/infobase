# inventory_app/views.py
from django.shortcuts import render, get_object_or_404
from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_http_methods
from django.utils import timezone
from .models import InventoryItem, Location
from finder.models import Remains
import json
import logging
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from common.utils.access_mixin import UserGroupRequiredMixin
from django.views.generic import TemplateView
from .serializers import LocationSerializer

logger = logging.getLogger(__name__)



def export_report(request):
    """Экспорт отчета в Excel с группировкой по родительским позициям"""
    
    # Получаем все родительские позиции
    parent_items = InventoryItem.objects.filter(parent__isnull=True).order_by('created_at')
    
    # Создаем workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Инвентаризация"
    
    # ===== СТИЛИ =====
    title_font = Font(bold=True, size=16, color="1F3A55")
    title_alignment = Alignment(horizontal="center", vertical="center")
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F3A55", end_color="1F3A55", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    cell_border = Border(
        left=Side(style='thin', color='D8DEE3'),
        right=Side(style='thin', color='D8DEE3'),
        top=Side(style='thin', color='D8DEE3'),
        bottom=Side(style='thin', color='D8DEE3')
    )
    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    parent_font = Font(bold=True, color="1F3A55", size=12)
    child_font = Font(color="1E3A8A")
    
    # ===== ЗАГОЛОВОК =====
    ws.merge_cells('A1:H1')
    title_cell = ws['A1']
    title_cell.value = "ОТЧЕТ ПО УЧЕТУ МАТЕРИАЛЬНЫХ ЦЕННОСТЕЙ"
    title_cell.font = title_font
    title_cell.alignment = title_alignment
    
    ws.merge_cells('A2:H2')
    date_cell = ws['A2']
    date_cell.value = f"Дата формирования: {timezone.now().strftime('%d.%m.%Y %H:%M')}"
    date_cell.font = Font(size=10, color="666666")
    date_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # ===== ЗАГОЛОВКИ ТАБЛИЦЫ =====
    headers = [
        '№', '№ ТН, ТТН', 'Дата поступления', 
        'Артикул', 'Наименование', 'Место хранения', 'Кол-во', 'Комментарий'
    ]
    
    header_row = 4
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = cell_border
    
    # ===== ЗАПИСЬ ДАННЫХ =====
    row_num = header_row + 1
    counter = 1
    
    def write_item(item, level=0):
        nonlocal row_num, counter
        
        level_prefix = '  ' * level
        
        # Исправлено: получаем название локации, а не объект
        location_name = item.location.name if item.location else ''
        
        data = [
            counter if level == 0 else '',
            item.number or '',
            item.arrival_date.strftime('%d.%m.%Y') if item.arrival_date else '',
            item.article or '',
            f"{level_prefix}{item.name or ''}",
            location_name,  # Теперь это строка, а не объект Location
            item.quantity or 0,
            item.comment or ''
        ]
        
        current_row = row_num
        
        for col, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = cell_border
            cell.alignment = cell_alignment
            
            if level == 0:
                cell.font = parent_font
                if col == 1:
                    cell.alignment = center_alignment
            elif level == 1:
                cell.font = child_font
                if col == 5:
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
            else:
                if col == 5:
                    cell.alignment = Alignment(horizontal="left", vertical="center", indent=level, wrap_text=True)
        
        counter += 1
        row_num += 1
        
        if level == 0:
            block_children = []
        
        children = item.children.all().order_by('created_at')
        for child in children:
            child_row = write_item(child, level + 1)
            if level == 0:
                block_children.append(child_row)
        
        if level == 0 and children.exists():
            block_end = row_num - 1 if block_children else current_row
            
            if block_children:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=block_end, end_column=1)
                ws.cell(row=current_row, column=1).alignment = center_alignment
                ws.cell(row=current_row, column=1).value = counter - len(children) - 1
            
            for r in range(current_row, block_end + 1):
                for c in range(1, 9):
                    cell = ws.cell(row=r, column=c)
                    if r == current_row:
                        cell.border = Border(
                            left=Side(style='medium', color='1F3A55'),
                            right=Side(style='medium', color='1F3A55'),
                            top=Side(style='medium', color='1F3A55'),
                            bottom=Side(style='thin', color='1F3A55') if r < block_end else Side(style='medium', color='1F3A55')
                        )
                    elif r == block_end:
                        cell.border = Border(
                            left=Side(style='medium', color='1F3A55'),
                            right=Side(style='medium', color='1F3A55'),
                            top=Side(style='thin', color='1F3A55'),
                            bottom=Side(style='medium', color='1F3A55')
                        )
                    else:
                        cell.border = Border(
                            left=Side(style='medium', color='1F3A55'),
                            right=Side(style='medium', color='1F3A55'),
                            top=Side(style='thin', color='1F3A55'),
                            bottom=Side(style='thin', color='1F3A55')
                        )
        
        return row_num
    
    for item in parent_items:
        row_num = write_item(item, 0)
    
    # ===== НАСТРОЙКА ШИРИНЫ КОЛОНОК =====
    column_widths = {
        'A': 5, 'B': 15, 'C': 15, 'D': 18, 'E': 45, 'F': 25, 'G': 10, 'H': 30
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # ===== ИТОГИ =====
    total_row = row_num + 2
    ws.merge_cells(f'A{total_row}:G{total_row}')
    total_label = ws.cell(row=total_row, column=1, value="ИТОГО ПОЗИЦИЙ:")
    total_label.font = Font(bold=True, size=12, color="1F3A55")
    total_label.alignment = Alignment(horizontal="right", vertical="center")
    
    total_items = InventoryItem.objects.count()
    total_cell = ws.cell(row=total_row, column=8, value=total_items)
    total_cell.font = Font(bold=True, size=12, color="F64E11")
    total_cell.alignment = Alignment(horizontal="center", vertical="center")
    total_cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    
    total_quantity_row = total_row + 1
    ws.merge_cells(f'A{total_quantity_row}:G{total_quantity_row}')
    total_qty_label = ws.cell(row=total_quantity_row, column=1, value="ОБЩЕЕ КОЛИЧЕСТВО:")
    total_qty_label.font = Font(bold=True, size=12, color="1F3A55")
    total_qty_label.alignment = Alignment(horizontal="right", vertical="center")
    
    all_items = InventoryItem.objects.all()
    total_quantity = sum(item.quantity or 0 for item in all_items)
    total_qty_cell = ws.cell(row=total_quantity_row, column=8, value=total_quantity)
    total_qty_cell.font = Font(bold=True, size=12, color="F64E11")
    total_qty_cell.alignment = Alignment(horizontal="center", vertical="center")
    total_qty_cell.fill = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    
    # ===== ФОРМАТИРОВАНИЕ =====
    for row in range(1, total_quantity_row + 1):
        for col in range(1, 9):
            cell = ws.cell(row=row, column=col)
            if not cell.alignment:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    
    # Создаем ответ
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"inventory_report_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    wb.save(response)
    return response


class TableView(UserGroupRequiredMixin, TemplateView):
    """Главная страница с таблицей"""
    template_name = 'cars/index.html'
    group_required = ['sklad', 'everyone']
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        user = self.request.user
        user_group = None
        allowed_link_names = []

        if user.is_authenticated:
            user_groups = user.groups.all()
            for group in user_groups:
                user_group = group.name
                break
            
            try:
                from finder.models import LinkAccess
                allowed_links = LinkAccess.objects.filter(group__in=user_groups)
                allowed_link_names = [link.link_name for link in allowed_links]
            except ImportError:
                pass

        items = InventoryItem.objects.filter(parent__isnull=True).order_by('created_at')
        
        context['items'] = items
        context['user_group'] = user_group
        context['allowed_link_names'] = json.dumps(allowed_link_names)
        
        return context


# ===== API ДЛЯ ПОЛУЧЕНИЯ ЛОКАЦИЙ =====

def get_locations_api(request):
    """API для получения списка активных мест хранения"""
    try:
        locations = Location.objects.filter(is_active=True).order_by('name')
        data = [{'id': loc.id, 'name': loc.name} for loc in locations]
        return JsonResponse({'success': True, 'locations': data})
    except Exception as e:
        logger.error(f"Error getting locations: {e}")
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


# ===== API ДЛЯ СОХРАНЕНИЯ =====

@csrf_exempt
@require_http_methods(["POST"])
def save_row(request):
    """Сохранение новой строки"""
    try:
        data = json.loads(request.body)
        
        # Обработка локации
        location_id = data.get('location')
        location = None
        if location_id and location_id != '':
            try:
                location = Location.objects.get(id=int(location_id), is_active=True)
            except (Location.DoesNotExist, ValueError, TypeError):
                pass
        
        item = InventoryItem.objects.create(
            number=data.get('number', ''),
            arrival_date=data.get('arrival_date') or None,
            article=data.get('article', ''),
            name=data.get('name', ''),
            location=location,
            quantity=data.get('quantity'),
            comment=data.get('comment', ''),
            parent_id=data.get('parent') or None
        )
        
        return JsonResponse({
            'success': True,
            'id': item.id,
            'message': 'Строка успешно сохранена'
        })
        
    except Exception as e:
        logger.error(f"Error saving row: {e}")
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=400)


@csrf_exempt
@require_http_methods(["POST"])
def update_row(request, item_id):
    """Обновление существующей строки"""
    try:
        item = get_object_or_404(InventoryItem, id=item_id)
        data = json.loads(request.body)
        
        # Обработка локации - исправлено
        location_id = data.get('location')
        location = None
        if location_id and location_id != '' and location_id != 'null':
            try:
                location = Location.objects.get(id=int(location_id), is_active=True)
            except (Location.DoesNotExist, ValueError, TypeError):
                pass
        
        item.number = data.get('number', item.number)
        item.arrival_date = data.get('arrival_date') or item.arrival_date
        item.article = data.get('article', item.article)
        item.name = data.get('name', item.name)
        item.location = location  # Теперь точно Location или None
        item.quantity = data.get('quantity', item.quantity)
        item.comment = data.get('comment', item.comment)
        item.save()
        
        return JsonResponse({
            'success': True,
            'id': item.id,
            'message': 'Строка обновлена'
        })
        
    except Exception as e:
        logger.error(f"Error updating row: {e}")
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@csrf_exempt
@require_http_methods(["DELETE"])
def delete_row(request, item_id):
    """Удаление строки и всех её детей"""
    try:
        item = get_object_or_404(InventoryItem, id=item_id)
        item.delete()
        
        return JsonResponse({
            'success': True,
            'message': 'Строка удалена'
        })
        
    except Exception as e:
        logger.error(f"Error deleting row: {e}")
        return JsonResponse({
            'success': False,
            'error': str(e)
        }, status=500)


@csrf_exempt
@require_http_methods(["POST"])
def add_child(request):
    """Добавление дочерней позиции"""
    try:
        data = json.loads(request.body)
        parent_id = data.get('parent_id')
        article = data.get('article')
        name = data.get('name', article)
        quantity = data.get('quantity', 1)
        
        parent = InventoryItem.objects.get(id=parent_id)
        
        child = InventoryItem.objects.create(
            number=parent.number,
            arrival_date=parent.arrival_date,
            location=parent.location,  # Наследуем локацию от родителя
            article=article,
            name=name,
            quantity=quantity,
            comment=data.get('comment', ''),
            parent=parent
        )
        
        return JsonResponse({
            'success': True,
            'id': child.id,
            'message': 'Дочерняя позиция добавлена'
        })
    except InventoryItem.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Родительская позиция не найдена'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=400)
